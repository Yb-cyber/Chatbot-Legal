import asyncio
import io
import json
import os
import uuid
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from lightrag import QueryParam
from rag_manager import close_rag, get_config, get_rag, init_rag

from config import DATA_DIR, DATA_CATEGORIES

load_dotenv()

UPLOAD_DIR = Path("./uploads")
UPLOAD_DIR.mkdir(exist_ok=True)


@asynccontextmanager
async def lifespan(app: FastAPI):
    await init_rag() 
    cfg = get_config()
    print(f"RAG initialized: {cfg['llm_model']} + {cfg['embedding_model']}")
    yield
    await close_rag()


app = FastAPI(
    title="Freeal RAG API",
    description="RAG legal para freelancers (Argentina) — 100% OpenRouter gratis",
    version="2.0.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


class QueryRequest(BaseModel):
    question: str
    mode: str = Field(default="hybrid", pattern="^(naive|local|global|hybrid|mix)$")
    top_k: int = 50
    only_need_context: bool = False
    conversation_history: List[Dict[str, str]] = Field(default_factory=list)


class QueryResponse(BaseModel):
    answer: str
    context: Optional[str] = None
    mode: str


class UploadResponse(BaseModel):
    filename: str
    status: str
    message: str


class HealthResponse(BaseModel):
    status: str
    rag_initialized: bool
    working_dir: str
    llm_model: str
    embedding_model: str


class StatsResponse(BaseModel):
    working_dir: str
    files_in_working_dir: int
    upload_files: int

"""funcion para extraer a texto limpio los datos de pdf,docx,pptx,txt ya descargados """
def _extract_upload_text(filename: str, raw: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in (".txt", ".md"):
        return raw.decode("utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if suffix == ".docx":
        import docx

        doc = docx.Document(io.BytesIO(raw))
        return "\n".join(p.text for p in doc.paragraphs)
    if suffix == ".pptx":
        from pptx import Presentation

        prs = Presentation(io.BytesIO(raw))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)
        return "\n".join(parts)
    raise ValueError(f"Formato no soportado: {suffix}")

"""funcion para fragmnetar e indexar los datos,vectorizacion"""
@app.post("/upload", response_model=UploadResponse)
async def upload_document(file: UploadFile = File(...)):
    rag = get_rag()
    if not rag:
        raise HTTPException(status_code=503, detail="RAG no inicializado")

    allowed_extensions = {".txt", ".md", ".pdf", ".docx", ".pptx"}
    file_ext = Path(file.filename or "").suffix.lower()
    if file_ext not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail=f"Tipo no permitido. Permitidos: {sorted(allowed_extensions)}",
        )

    document_id = str(uuid.uuid4())
    file_path = UPLOAD_DIR / f"{document_id}_{file.filename}"
    try:
        raw = await file.read()
        file_path.write_bytes(raw)
        content = _extract_upload_text(file.filename, raw)
        if len(content.strip()) < 50:
            raise ValueError("El documento no contiene texto extraíble suficiente")
        await rag.ainsert(content)
        return UploadResponse(
            filename=file.filename,
            status="indexed",
            message=f"Documento indexado con ID: {document_id}",
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo indexar: {e}")


def _friendly_error(e: Exception) -> str:
    msg = str(e)
    if "401" in msg:
        return "OPENROUTER_API_KEY inválida. Revisa tu key en https://openrouter.ai/keys"
    if any(c in msg for c in ("429", "503", "504", "UNAVAILABLE")):
        return (
            "El modelo gratis está saturado o con rate limit. "
            "Reintenta en unos minutos."
        )
    if "404" in msg and "train" in msg.lower():
        return (
            "Tu cuenta de OpenRouter bloquea modelos gratis "
            "(privacy settings). Activa el uso de modelos free."
        )
    return f"Falló la consulta: {msg}"

"""llamada a la API , busqueda en archivos fragmentados,etc"""
@app.post("/query", response_model=QueryResponse)
async def query(request: QueryRequest):
    rag = get_rag()
    if not rag:
        raise HTTPException(status_code=503, detail="RAG no inicializado")

    try:
        param = QueryParam(
            mode=request.mode,
            top_k=request.top_k,
            only_need_context=request.only_need_context,
            conversation_history=request.conversation_history[-10:],
        )
        answer = await rag.aquery(request.question, param=param)
        if answer is None:
            return QueryResponse(
                answer=(
                    "No se pudo generar una respuesta. "
                    "Probablemente aún no hay documentos indexados: "
                    "ejecuta POST /update-data primero."
                ),
                context=None,
                mode=request.mode,
            )
        ctx = str(answer) if request.only_need_context else None
        return QueryResponse(answer=str(answer), context=ctx, mode=request.mode)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=_friendly_error(e))


@app.get("/health", response_model=HealthResponse)
async def health():
    cfg = get_config()
    working_path = Path(cfg["working_dir"])
    files_count = len(list(working_path.glob("*"))) if working_path.exists() else 0
    upload_files = len(list(UPLOAD_DIR.glob("*"))) if UPLOAD_DIR.exists() else 0
    return HealthResponse(
        status="healthy",
        rag_initialized=get_rag() is not None,
        working_dir=cfg["working_dir"],
        llm_model=cfg["llm_model"],
        embedding_model=cfg["embedding_model"],
    )


@app.get("/stats", response_model=StatsResponse)
async def stats():
    cfg = get_config()
    working_path = Path(cfg["working_dir"])
    files_count = len(list(working_path.glob("*"))) if working_path.exists() else 0
    upload_files = len(list(UPLOAD_DIR.glob("*"))) if UPLOAD_DIR.exists() else 0
    return StatsResponse(
        working_dir=cfg["working_dir"],
        files_in_working_dir=files_count,
        upload_files=upload_files,
    )


@app.delete("/documents/{doc_id}")
async def delete_document(doc_id: str):
    rag = get_rag()
    if not rag:
        raise HTTPException(status_code=503, detail="RAG no inicializado")
    try:
        await rag.delete_by_doc_id(doc_id)
        return {"status": "deleted", "doc_id": doc_id}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo eliminar: {e}")


class UpdateDataResponse(BaseModel):
    status: str
    message: str
    last_update: Optional[str] = None
    total_files: int = 0
    categories: Dict[str, int] = {}


class DataStatusResponse(BaseModel):
    last_update: Optional[str] = None
    total_files: int = 0
    categories: Dict[str, int] = {}
    metadata: Dict[str, Any] = {}


@app.post("/update-data", response_model=UpdateDataResponse)
async def update_data(full: bool = False, test: bool = False, minimal: bool = False):
    from download_argentina_legal import download_only
    from index_data import index_all_data
    import os

    rag = get_rag()
    if not rag:
        raise HTTPException(status_code=503, detail="RAG no inicializado")
    try:
        loop = asyncio.get_running_loop()

        # Paso 1: descarga en thread (sin asyncio.run)
        metadata = await loop.run_in_executor(
            None,
            lambda: download_only(full=full, test_mode=test, minimal=minimal),
        )

        # Paso 2: indexación en el event loop principal
        max_files = int(os.getenv("INDEX_MAX_FILES", "100"))
        index_stats = await index_all_data(
            rag, full=full, max_files=max_files
        )
        metadata["indexing"] = index_stats
        from download_argentina_legal import save_metadata
        save_metadata(metadata)

        categories_count = {}
        for category_dir in DATA_CATEGORIES.values():
            if category_dir.exists():
                categories_count[category_dir.name] = len(
                    list(category_dir.glob("*"))
                )
        return UpdateDataResponse(
            status="success",
            message="Datos actualizados e indexados",
            last_update=metadata.get("last_update"),
            total_files=len(metadata.get("downloads", {})),
            categories=categories_count,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo actualizar: {e}")


@app.get("/data-status", response_model=DataStatusResponse)
async def data_status():
    metadata_file = DATA_DIR / "metadata.json"
    if not metadata_file.exists():
        return DataStatusResponse(
            last_update=None, total_files=0, categories={}, metadata={}
        )
    try:
        with open(metadata_file, "r", encoding="utf-8") as f:
            metadata = json.load(f)
        categories_count = {}
        for category_dir in DATA_CATEGORIES.values():
            if category_dir.exists():
                categories_count[category_dir.name] = len(
                    list(category_dir.glob("*"))
                )
        return DataStatusResponse(
            last_update=metadata.get("last_update"),
            total_files=len(metadata.get("downloads", {})),
            categories=categories_count,
            metadata=metadata,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo leer metadata: {e}")


@app.get("/data/{category}")
async def list_category_files(category: str):
    if category not in DATA_CATEGORIES:
        raise HTTPException(
            status_code=400,
            detail=f"Categoría inválida. Válidas: {list(DATA_CATEGORIES.keys())}",
        )
    category_dir = DATA_CATEGORIES[category]
    if not category_dir.exists():
        return {"category": category, "files": []}
    files = []
    for file_path in category_dir.glob("*"):
        if file_path.is_file():
            files.append(
                {
                    "name": file_path.name,
                    "size": file_path.stat().st_size,
                    "modified": datetime.fromtimestamp(
                        file_path.stat().st_mtime
                    ).isoformat(),
                }
            )
    return {"category": category, "files": files}


# Servir frontend estático
FRONTEND_DIR = Path(__file__).parent / "frontend"
if FRONTEND_DIR.exists():
    app.mount("/", StaticFiles(directory=str(FRONTEND_DIR), html=True), name="frontend")


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
